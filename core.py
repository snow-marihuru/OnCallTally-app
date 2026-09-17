"""電話当番の集計ロジック。

Web UI (app.py) から独立させており、将来ローカル運用に切り替える際は
このモジュールだけを利用すればよい構成にしている。

祝日判定には jpholiday を使用している。jpholiday は祝日データを
ライブラリ内部に保持しているため、法改正等に対応するには
年1回程度ライブラリ本体をアップデートすること。
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO

import jpholiday
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CATEGORIES = ("weekday", "saturday", "sunday_holiday")
CATEGORY_LABELS = {
    "weekday": "平日日数",
    "saturday": "土曜日数",
    "sunday_holiday": "日祝日日数",
}
CATEGORY_SHORT_LABELS = {
    "weekday": "平日",
    "saturday": "土曜",
    "sunday_holiday": "日祝日",
}
ALLOWANCE_RATES = {
    "weekday": 600,
    "saturday": 900,
    "sunday_holiday": 1200,
}


@dataclass
class Assignment:
    name: str
    start: date
    end: date


def classify_day(day: date) -> str:
    """日付を 'weekday' / 'saturday' / 'sunday_holiday' に分類する（自動判定）。"""
    if day.weekday() == 6 or jpholiday.is_holiday(day):
        return "sunday_holiday"
    if day.weekday() == 5:
        return "saturday"
    return "weekday"


def resolve_category(day: date, overrides: dict[str, str] | None = None) -> str:
    """日付の区分を決定する。overrides に手動指定があればそちらを優先する。

    会社の公休日や、祝日だが平日扱いにする日など、自動判定（classify_day）と
    実態がずれるケースを人手で補正できるようにするための仕組み。overrides は
    ISO日付文字列("YYYY-MM-DD") -> 区分名 の辞書。
    """
    if overrides:
        override = overrides.get(day.isoformat())
        if override in CATEGORIES:
            return override
    return classify_day(day)


def aggregate(
    assignments: list[Assignment],
    period_start: date,
    period_end: date,
    overrides: dict[str, str] | None = None,
) -> dict[str, dict[str, int]]:
    """当番期間の一覧から、集計期間内における人ごとの区分別日数を集計する。

    担当期間が集計期間からはみ出している場合は、集計期間内の日数のみを数える。
    overrides が指定された日は、自動判定ではなく overrides の区分でカウントする。
    """
    result: dict[str, dict[str, int]] = {}

    for assignment in assignments:
        clipped_start = max(assignment.start, period_start)
        clipped_end = min(assignment.end, period_end)
        if clipped_start > clipped_end:
            continue

        counts = result.setdefault(assignment.name, {c: 0 for c in CATEGORIES})

        day = clipped_start
        while day <= clipped_end:
            counts[resolve_category(day, overrides)] += 1
            day += timedelta(days=1)

    return result


def find_overlapping_days(
    assignments: list[Assignment], period_start: date, period_end: date
) -> dict[date, list[str]]:
    """集計期間内で、同じ日に複数の当番期間が重なっている日を検出する。

    同一人物の重複登録（同じ日を含む期間の二重登録）も、日数の二重カウントに
    つながるため検出対象に含める。戻り値は重複がある日付のみを含む。
    """
    day_owners: dict[date, list[str]] = {}

    for assignment in assignments:
        clipped_start = max(assignment.start, period_start)
        clipped_end = min(assignment.end, period_end)
        if clipped_start > clipped_end:
            continue

        day = clipped_start
        while day <= clipped_end:
            day_owners.setdefault(day, []).append(assignment.name)
            day += timedelta(days=1)

    return {day: names for day, names in sorted(day_owners.items()) if len(names) > 1}


def find_uncovered_days(
    assignments: list[Assignment], period_start: date, period_end: date
) -> list[date]:
    """集計期間内で、どの当番期間にも含まれていない日(空白日)を検出する。"""
    covered: set[date] = set()

    for assignment in assignments:
        clipped_start = max(assignment.start, period_start)
        clipped_end = min(assignment.end, period_end)
        if clipped_start > clipped_end:
            continue

        day = clipped_start
        while day <= clipped_end:
            covered.add(day)
            day += timedelta(days=1)

    uncovered = []
    day = period_start
    while day <= period_end:
        if day not in covered:
            uncovered.append(day)
        day += timedelta(days=1)

    return uncovered


def build_calendar(
    assignments: list[Assignment],
    period_start: date,
    period_end: date,
    overrides: dict[str, str] | None = None,
) -> list[list[dict]]:
    """集計期間全体を1つの連続したカレンダー(週単位の行)として組み立てる。

    月をまたぐ集計期間（例: 7/16〜8/15）でも、月ごとに表を分けず1つの表で
    表示できるよう、period_start を含む週の月曜日から period_end を含む週の
    日曜日までを週単位（月曜始まり）で並べる。集計期間外の日（前後の端数週の
    日）も in_period=False として含め、グレー表示などに使えるようにする。

    各セルには当日の区分(weekday/saturday/sunday_holiday。overridesがあれば
    そちらを優先)と、当番の担当者名一覧を持たせる。
    """
    day_names: dict[date, list[str]] = {}
    for assignment in assignments:
        clipped_start = max(assignment.start, period_start)
        clipped_end = min(assignment.end, period_end)
        if clipped_start > clipped_end:
            continue
        day = clipped_start
        while day <= clipped_end:
            day_names.setdefault(day, []).append(assignment.name)
            day += timedelta(days=1)

    first_monday = period_start - timedelta(days=period_start.weekday())
    last_sunday = period_end + timedelta(days=(6 - period_end.weekday()) % 7)

    weeks: list[list[dict]] = []
    day = first_monday
    while day <= last_sunday:
        week_cells = []
        for _ in range(7):
            in_period = period_start <= day <= period_end
            week_cells.append(
                {
                    "date": day,
                    "in_period": in_period,
                    "category": resolve_category(day, overrides) if in_period else None,
                    "overridden": bool(in_period and overrides and day.isoformat() in overrides),
                    "names": day_names.get(day, []) if in_period else [],
                }
            )
            day += timedelta(days=1)
        weeks.append(week_cells)

    return weeks


def calculate_allowance(counts: dict[str, int]) -> int:
    """区分別の当番日数から手当支給額(円)を計算する。

    平日600円/土曜900円/日祝日1200円の単価で、日数に応じて支給される想定。
    """
    return sum(counts[c] * ALLOWANCE_RATES[c] for c in CATEGORIES)


def build_excel(
    aggregation: dict[str, dict[str, int]], period_start: date, period_end: date
) -> BytesIO:
    """集計結果からExcelファイル(.xlsx)を生成し、BytesIOで返す。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "当番集計"

    ws["A1"] = f"集計期間: {period_start.isoformat()} 〜 {period_end.isoformat()}"
    ws["A1"].font = Font(bold=True)

    header_row = 3
    headers = ["担当者", *[CATEGORY_LABELS[c] for c in CATEGORIES], "合計日数", "手当支給額(円)"]
    total_days_col = len(headers) - 1
    allowance_col = len(headers)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    row = header_row + 1
    for name in sorted(aggregation.keys()):
        counts = aggregation[name]
        total = sum(counts[c] for c in CATEGORIES)
        ws.cell(row=row, column=1, value=name)
        for col, category in enumerate(CATEGORIES, start=2):
            ws.cell(row=row, column=col, value=counts[category])
        ws.cell(row=row, column=total_days_col, value=total)
        ws.cell(row=row, column=allowance_col, value=calculate_allowance(counts))
        row += 1

    totals = {c: 0 for c in CATEGORIES}
    for counts in aggregation.values():
        for category in CATEGORIES:
            totals[category] += counts[category]
    grand_total = sum(totals.values())
    grand_allowance = calculate_allowance(totals)

    ws.cell(row=row, column=1, value="合計").font = Font(bold=True)
    for col, category in enumerate(CATEGORIES, start=2):
        ws.cell(row=row, column=col, value=totals[category]).font = Font(bold=True)
    ws.cell(row=row, column=total_days_col, value=grand_total).font = Font(bold=True)
    ws.cell(row=row, column=allowance_col, value=grand_allowance).font = Font(bold=True)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # 手当支給額の検証用内訳。区分別の手当合計を個別に示すことで、
    # 上表の「合計」行の手当支給額と一致するかを目視で確認できるようにする。
    breakdown_row = row + 2
    ws.cell(row=breakdown_row, column=1, value="手当内訳(検証用)").font = Font(bold=True)
    breakdown_row += 1
    for category in CATEGORIES:
        ws.cell(row=breakdown_row, column=1, value=f"{CATEGORY_SHORT_LABELS[category]}手当合計")
        ws.cell(row=breakdown_row, column=2, value=totals[category] * ALLOWANCE_RATES[category])
        breakdown_row += 1
    ws.cell(row=breakdown_row, column=1, value="手当支給額合計").font = Font(bold=True)
    ws.cell(row=breakdown_row, column=2, value=grand_allowance).font = Font(bold=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def default_period(today: date) -> tuple[date, date]:
    """給与締め(毎月15日)に合わせた、today が含まれる集計サイクルを返す。

    サイクルは「前月16日〜当月15日」。today が16日以降なら
    「当月16日〜翌月15日」の進行中サイクルを返す。
    """
    if today.day >= 16:
        start = date(today.year, today.month, 16)
        end_year, end_month = today.year, today.month + 1
        if end_month > 12:
            end_year, end_month = end_year + 1, 1
        end = date(end_year, end_month, 15)
    else:
        end = date(today.year, today.month, 15)
        start_year, start_month = today.year, today.month - 1
        if start_month < 1:
            start_year, start_month = start_year - 1, 12
        start = date(start_year, start_month, 16)
    return start, end
