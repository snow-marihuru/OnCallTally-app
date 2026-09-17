"""core.py の集計ロジックに対する自動テスト。

日付は固定値を使い、jpholiday のバージョンアップ等で祝日データが変わった場合に
テストが検出できるようにしている。2026年7月20日(海の日、月曜)を祝日の代表例に、
7/21〜7/24を祝日を含まない平日として使用している。
"""
from datetime import date

from openpyxl import load_workbook

import core


def test_classify_day_weekday():
    assert core.classify_day(date(2026, 7, 21)) == "weekday"


def test_classify_day_saturday():
    assert core.classify_day(date(2026, 7, 18)) == "saturday"


def test_classify_day_sunday():
    assert core.classify_day(date(2026, 7, 19)) == "sunday_holiday"


def test_classify_day_national_holiday_on_weekday():
    # 海の日(7月第3月曜日)。2026年は7/20。祝日は曜日に関わらず sunday_holiday 扱い。
    assert core.classify_day(date(2026, 7, 20)) == "sunday_holiday"


def test_resolve_category_without_override_matches_classify_day():
    day = date(2026, 7, 21)
    assert core.resolve_category(day) == core.classify_day(day) == "weekday"


def test_resolve_category_with_override():
    day = date(2026, 7, 21)  # 自動判定は weekday
    overrides = {day.isoformat(): "sunday_holiday"}
    assert core.resolve_category(day, overrides) == "sunday_holiday"


def test_resolve_category_ignores_unknown_category_value():
    day = date(2026, 7, 21)
    overrides = {day.isoformat(): "not_a_real_category"}
    assert core.resolve_category(day, overrides) == "weekday"


def test_aggregate_full_month_cycle():
    assignments = [core.Assignment("Aさん", date(2026, 7, 16), date(2026, 8, 15))]
    result = core.aggregate(assignments, date(2026, 7, 16), date(2026, 8, 15))
    assert result == {"Aさん": {"weekday": 20, "saturday": 5, "sunday_holiday": 6}}


def test_aggregate_clips_to_period():
    # 担当期間が集計期間の外にはみ出す分はカウントしない。
    assignments = [core.Assignment("Aさん", date(2026, 7, 1), date(2026, 7, 31))]
    result = core.aggregate(assignments, date(2026, 7, 16), date(2026, 7, 20))
    assert sum(result["Aさん"].values()) == 5


def test_aggregate_groups_same_person_across_multiple_assignments():
    assignments = [
        core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 22)),
        core.Assignment("Aさん", date(2026, 7, 23), date(2026, 7, 24)),
    ]
    result = core.aggregate(assignments, date(2026, 7, 21), date(2026, 7, 24))
    assert list(result.keys()) == ["Aさん"]
    assert result["Aさん"]["weekday"] == 4


def test_aggregate_applies_overrides():
    assignments = [core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 21))]
    period_start = period_end = date(2026, 7, 21)

    default_result = core.aggregate(assignments, period_start, period_end)
    assert default_result["Aさん"] == {"weekday": 1, "saturday": 0, "sunday_holiday": 0}

    overrides = {"2026-07-21": "sunday_holiday"}
    overridden_result = core.aggregate(assignments, period_start, period_end, overrides)
    assert overridden_result["Aさん"] == {"weekday": 0, "saturday": 0, "sunday_holiday": 1}


def test_find_overlapping_days_detects_different_people():
    assignments = [
        core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 22)),
        core.Assignment("Bさん", date(2026, 7, 22), date(2026, 7, 23)),
    ]
    overlaps = core.find_overlapping_days(assignments, date(2026, 7, 21), date(2026, 7, 23))
    assert overlaps == {date(2026, 7, 22): ["Aさん", "Bさん"]}


def test_find_overlapping_days_detects_same_person_double_booking():
    assignments = [
        core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 22)),
        core.Assignment("Aさん", date(2026, 7, 22), date(2026, 7, 23)),
    ]
    overlaps = core.find_overlapping_days(assignments, date(2026, 7, 21), date(2026, 7, 23))
    assert overlaps == {date(2026, 7, 22): ["Aさん", "Aさん"]}


def test_find_overlapping_days_no_overlap():
    assignments = [
        core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 22)),
        core.Assignment("Bさん", date(2026, 7, 23), date(2026, 7, 24)),
    ]
    overlaps = core.find_overlapping_days(assignments, date(2026, 7, 21), date(2026, 7, 24))
    assert overlaps == {}


def test_find_uncovered_days_detects_gap():
    assignments = [core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 22))]
    uncovered = core.find_uncovered_days(assignments, date(2026, 7, 21), date(2026, 7, 24))
    assert uncovered == [date(2026, 7, 23), date(2026, 7, 24)]


def test_find_uncovered_days_full_coverage():
    assignments = [core.Assignment("Aさん", date(2026, 7, 21), date(2026, 7, 24))]
    uncovered = core.find_uncovered_days(assignments, date(2026, 7, 21), date(2026, 7, 24))
    assert uncovered == []


def test_build_calendar_covers_whole_period_and_starts_on_monday():
    weeks = core.build_calendar([], date(2026, 7, 16), date(2026, 8, 15))
    # 月曜始まり: 各週の1列目は月曜、7列目は日曜。
    for week in weeks:
        assert week[0]["date"].weekday() == 0
        assert week[6]["date"].weekday() == 6

    in_period_dates = [c["date"] for w in weeks for c in w if c["in_period"]]
    assert len(in_period_dates) == 31
    assert in_period_dates[0] == date(2026, 7, 16)
    assert in_period_dates[-1] == date(2026, 8, 15)


def test_build_calendar_marks_overridden_cells():
    weeks = core.build_calendar(
        [], date(2026, 7, 21), date(2026, 7, 21), {"2026-07-21": "sunday_holiday"}
    )
    cell = next(c for w in weeks for c in w if c["in_period"])
    assert cell["category"] == "sunday_holiday"
    assert cell["overridden"] is True


def test_calculate_allowance():
    counts = {"weekday": 6, "saturday": 2, "sunday_holiday": 2}
    # 6*600 + 2*900 + 2*1200 = 3600 + 1800 + 2400 = 7800
    assert core.calculate_allowance(counts) == 7800


def test_build_excel_contains_totals_row_and_allowance():
    aggregation = {
        "Aさん": {"weekday": 6, "saturday": 2, "sunday_holiday": 2},
        "Bさん": {"weekday": 14, "saturday": 3, "sunday_holiday": 4},
    }
    buffer = core.build_excel(aggregation, date(2026, 7, 16), date(2026, 8, 15))
    wb = load_workbook(buffer)
    rows = list(wb.active.iter_rows(values_only=True))

    # Aさん: 6*600+2*900+2*1200=7800 / Bさん: 14*600+3*900+4*1200=15900
    assert rows[3] == ("Aさん", 6, 2, 2, 10, 7800)
    assert rows[4] == ("Bさん", 14, 3, 4, 21, 15900)
    assert rows[5] == ("合計", 20, 5, 6, 31, 23700)

    # 区分別の手当内訳(検証用) - 合計行の手当支給額(23700)と一致すること
    assert rows[7] == ("手当内訳(検証用)", None, None, None, None, None)
    assert rows[8] == ("平日手当合計", 12000, None, None, None, None)
    assert rows[9] == ("土曜手当合計", 4500, None, None, None, None)
    assert rows[10] == ("日祝日手当合計", 7200, None, None, None, None)
    assert rows[11] == ("手当支給額合計", 23700, None, None, None, None)


def test_default_period_before_15th_uses_previous_cycle():
    start, end = core.default_period(date(2026, 7, 10))
    assert (start, end) == (date(2026, 6, 16), date(2026, 7, 15))


def test_default_period_after_16th_uses_current_cycle():
    start, end = core.default_period(date(2026, 7, 26))
    assert (start, end) == (date(2026, 7, 16), date(2026, 8, 15))


def test_default_period_handles_year_boundary():
    start, end = core.default_period(date(2026, 1, 10))
    assert (start, end) == (date(2025, 12, 16), date(2026, 1, 15))

    start, end = core.default_period(date(2026, 12, 20))
    assert (start, end) == (date(2026, 12, 16), date(2027, 1, 15))
